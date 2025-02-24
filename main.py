#!/usr/bin/env python3
import os
import sys
import zipfile
import shutil
import logging
import tempfile
import re
from lxml import etree as LET

# Setup logging: output to file and console.
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.FileHandler("metadata_removal.log"), logging.StreamHandler()]
)

# --- PDF processing ---
from PyPDF2 import PdfReader, PdfWriter

# --- DOCX processing ---
from docx import Document

# --- XLSX processing ---
from openpyxl import load_workbook

# --- PPTX processing (if available) ---
try:
    from pptx import Presentation
    CAN_HANDLE_PPTX = True
except ImportError:
    CAN_HANDLE_PPTX = False

# --- Legacy DOC handling ---
try:
    import olefile
except ImportError:
    olefile = None

# --- PyQt5 GUI ---
from PyQt5.QtWidgets import QApplication, QWidget, QPushButton, QVBoxLayout, QFileDialog, QMessageBox, QLabel

################################################################
# XML Modification Helpers (using lxml)
################################################################
def modify_xml_in_zip_file(file_path, internal_path, modify_func):
    """
    Open a ZIP-based file (e.g. DOCX or XLSX) and modify the XML stored at internal_path.
    This function uses lxml to parse the XML while preserving the original namespace prefixes and XML declaration.
    It captures the original XML declaration (if any) and re-prepends it after modification.
    """
    temp_file = tempfile.NamedTemporaryFile(delete=False)
    temp_file.close()
    xml_decl_re = re.compile(br'^<\?xml\s+[^>]+\?>\s*')
    try:
        with zipfile.ZipFile(file_path, 'r') as zin, \
             zipfile.ZipFile(temp_file.name, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.lower() == internal_path.lower():
                    try:
                        m = xml_decl_re.match(data)
                        xml_decl = m.group(0) if m else b''
                        data_no_decl = data[m.end():] if m else data
                        # Parse with lxml; this preserves namespaces and structure.
                        parser = LET.XMLParser(remove_blank_text=False)
                        root = LET.fromstring(data_no_decl, parser)
                        modify_func(root)
                        modified = LET.tostring(root, encoding="UTF-8", pretty_print=False)
                        # Re-prepend the original XML declaration if it existed.
                        if xml_decl:
                            data = xml_decl + modified
                        else:
                            data = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + modified
                        logging.info(f"Modified {internal_path} in {file_path}")
                    except Exception as e:
                        logging.exception(f"Failed to modify {internal_path} in {file_path}: {e}")
                zout.writestr(item, data)
        shutil.move(temp_file.name, file_path)
        return True
    except Exception as e:
        logging.exception(f"Error modifying {internal_path} in {file_path}: {e}")
        if os.path.exists(temp_file.name):
            os.remove(temp_file.name)
        return False

################################################################
# XML Clear Functions for Confidential Fields
################################################################
def clear_confidential_core_fields(root):
    """
    In docProps/core.xml, clear confidential fields.
    For this example, always clear the "creator" field.
    (Comparison is done on the local tag name, case-insensitively.)
    """
    for elem in root.iter():
        local = elem.tag.split('}', 1)[-1] if '}' in elem.tag else elem.tag
        if local.lower() == "creator":
            elem.text = ""
            logging.info("Cleared core field: creator")

def clear_confidential_app_fields(root):
    """
    In docProps/app.xml, clear confidential fields.
    For this example, always clear the "Company" field.
    """
    for elem in root.iter():
        local = elem.tag.split('}', 1)[-1] if '}' in elem.tag else elem.tag
        if local.lower() == "company":
            elem.text = ""
            logging.info("Cleared app field: Company")

def clear_all_custom_fields(root):
    """
    In docProps/custom.xml, remove all child elements so that the file becomes empty.
    The root element and its namespaces are preserved.
    """
    for child in list(root):
        root.remove(child)
    logging.info("Cleared all custom properties in custom.xml")

################################################################
# Metadata Removal Functions for Each File Type
################################################################
def remove_metadata_from_pdf(pdf_path):
    """Remove PDF metadata by rewriting pages with empty metadata."""
    try:
        logging.info(f"Processing PDF: {pdf_path}")
        reader = PdfReader(pdf_path)
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.add_metadata({
            "/Title": "",
            "/Author": "",
            "/Subject": "",
            "/Creator": "",
            "/Producer": "",
            "/CreationDate": "",
            "/ModDate": ""
        })
        with open(pdf_path, 'wb') as f:
            writer.write(f)
        return True
    except Exception as e:
        logging.exception(f"Error processing PDF {pdf_path}: {e}")
        return False

def remove_metadata_from_docx(docx_path):
    """Remove DOCX metadata by clearing core properties and modifying XML parts."""
    try:
        logging.info(f"Processing DOCX: {docx_path}")
        doc = Document(docx_path)
        for field in ['author', 'comments', 'category', 'content_status',
                      'identifier', 'keywords', 'language', 'last_modified_by',
                      'revision', 'subject', 'title', 'version']:
            try:
                setattr(doc.core_properties, field, "")
            except Exception:
                pass
        doc.settings.odd_and_even_pages_header_footer = False
        doc.save(docx_path)
        # Modify docProps/core.xml to clear confidential fields.
        modify_xml_in_zip_file(docx_path, "docProps/core.xml", clear_confidential_core_fields)
        # Modify docProps/app.xml to clear the Company field.
        modify_xml_in_zip_file(docx_path, "docProps/app.xml", clear_confidential_app_fields)
        # Modify docProps/custom.xml to remove all custom properties.
        modify_xml_in_zip_file(docx_path, "docProps/custom.xml", clear_all_custom_fields)
        return True
    except Exception as e:
        logging.exception(f"Error processing DOCX {docx_path}: {e}")
        return False

def remove_metadata_from_xlsx(xlsx_path):
    """Remove XLSX metadata by clearing workbook properties and modifying XML parts."""
    try:
        logging.info(f"Processing XLSX: {xlsx_path}")
        wb = load_workbook(xlsx_path)
        for field in ['creator', 'title', 'subject', 'description',
                      'keywords', 'category', 'comments', 'last_modified_by',
                      'company', 'manager']:
            try:
                setattr(wb.properties, field, "")
            except Exception:
                pass
        wb.save(xlsx_path)
        modify_xml_in_zip_file(xlsx_path, "docProps/core.xml", clear_confidential_core_fields)
        modify_xml_in_zip_file(xlsx_path, "docProps/app.xml", clear_confidential_app_fields)
        modify_xml_in_zip_file(xlsx_path, "docProps/custom.xml", clear_all_custom_fields)
        return True
    except Exception as e:
        logging.exception(f"Error processing XLSX {xlsx_path}: {e}")
        return False

def remove_metadata_from_doc(doc_path):
    """For legacy DOC files, metadata removal is not implemented."""
    logging.warning(f"Metadata removal for DOC is not implemented: {doc_path}")
    return True

def remove_metadata_from_pptx(pptx_path):
    """Remove PPTX metadata using python-pptx."""
    if not CAN_HANDLE_PPTX:
        logging.warning("python-pptx not installed; cannot process PPTX.")
        return False
    try:
        logging.info(f"Processing PPTX: {pptx_path}")
        ppt = Presentation(pptx_path)
        props = ppt.core_properties
        for field in ['author', 'category', 'comments', 'content_status',
                      'created', 'identifier', 'keywords', 'last_modified_by',
                      'last_printed', 'modified', 'revision', 'subject', 'title']:
            try:
                setattr(props, field, "" if getattr(props, field) is not None else None)
            except Exception:
                pass
        ppt.save(pptx_path)
        return True
    except Exception as e:
        logging.exception(f"Error processing PPTX {pptx_path}: {e}")
        return False

def remove_metadata_from_ppt(ppt_path):
    """For legacy PPT files, metadata removal is not implemented."""
    logging.warning(f"Metadata removal for PPT is not implemented: {ppt_path}")
    return True

def remove_metadata(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return remove_metadata_from_pdf(file_path)
    elif ext == '.docx':
        return remove_metadata_from_docx(file_path)
    elif ext in ['.xlsx', '.xls']:
        return remove_metadata_from_xlsx(file_path)
    elif ext == '.doc':
        return remove_metadata_from_doc(file_path)
    elif ext == '.pptx':
        return remove_metadata_from_pptx(file_path)
    elif ext == '.ppt':
        return remove_metadata_from_ppt(file_path)
    else:
        logging.warning(f"Unsupported file type: {file_path}")
        return False

################################################################
# Folder Processing Function
################################################################
def process_folder(input_folder, output_folder):
    processed = 0
    failures = 0
    for root, dirs, files in os.walk(input_folder):
        rel = os.path.relpath(root, input_folder)
        target = os.path.join(output_folder, rel)
        os.makedirs(target, exist_ok=True)
        for file in files:
            ext = os.path.splitext(file)[1].lower()
            if ext in ['.pdf', '.docx', '.xlsx', '.xls', '.pptx', '.ppt', '.doc']:
                original = os.path.join(root, file)
                new_path = os.path.join(target, file)
                try:
                    shutil.copy2(original, new_path)
                    logging.info(f"Copied: {new_path}")
                except Exception as e:
                    logging.exception(f"Error copying {original}: {e}")
                    failures += 1
                    continue
                if remove_metadata(new_path):
                    processed += 1
                else:
                    failures += 1
    return processed, failures

################################################################
# PyQt5 GUI Application
################################################################
class MetadataRemoverGUI(QWidget):
    def __init__(self):
        super().__init__()
        self.initUI()
    
    def initUI(self):
        self.setWindowTitle("Metadata Remover")
        self.resize(400, 200)
        layout = QVBoxLayout()
        
        self.infoLabel = QLabel("Select a folder to remove metadata:")
        layout.addWidget(self.infoLabel)
        
        self.selectButton = QPushButton("Select Folder")
        self.selectButton.clicked.connect(self.select_folder)
        layout.addWidget(self.selectButton)
        
        self.statusLabel = QLabel("")
        layout.addWidget(self.statusLabel)
        
        self.setLayout(layout)
    
    def select_folder(self):
        folder = QFileDialog.getExistingDirectory(self, "Select Folder")
        if folder:
            self.infoLabel.setText(f"Selected: {folder}")
            parent = os.path.dirname(folder)
            base = os.path.basename(folder)
            output_folder = os.path.join(parent, base + "_no_metadata")
            os.makedirs(output_folder, exist_ok=True)
            processed, failures = process_folder(folder, output_folder)
            self.statusLabel.setText(f"Processed: {processed} files, Failures: {failures}")
            QMessageBox.information(self, "Done",
                                    f"Processing complete.\nProcessed: {processed} files.\nFailures: {failures}\nOutput: {output_folder}")

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MetadataRemoverGUI()
    window.show()
    sys.exit(app.exec_())
